from flask import Flask, render_template, request, jsonify, session, redirect, url_for
from pathlib import Path
from datetime import datetime
import sqlite3, hashlib, secrets, json, os, re, shutil

BASE_DIR = Path(__file__).resolve().parent
SEED_DB_PATH = BASE_DIR / "database" / "medsafe.db"
DB_PATH = Path(os.environ.get("MEDSAFE_DB_PATH", str(SEED_DB_PATH)))

app = Flask(__name__)
app.secret_key = os.environ.get("MEDSAFE_SECRET_KEY", secrets.token_hex(32))
app.config["MAX_CONTENT_LENGTH"] = 20 * 1024 * 1024

def db():
    DB_PATH.parent.mkdir(parents=True, exist_ok=True)
    if DB_PATH != SEED_DB_PATH and not DB_PATH.exists() and SEED_DB_PATH.exists():
        shutil.copy2(SEED_DB_PATH, DB_PATH)
    c = sqlite3.connect(DB_PATH)
    c.row_factory = sqlite3.Row
    c.execute("PRAGMA foreign_keys=ON")
    return c

def init_db():
    c = db()
    c.executescript("""
    CREATE TABLE IF NOT EXISTS hospitals(
      id INTEGER PRIMARY KEY AUTOINCREMENT, hospital_id TEXT UNIQUE NOT NULL,
      hospital_name TEXT NOT NULL, email TEXT UNIQUE NOT NULL, phone TEXT, address TEXT, created_at TEXT NOT NULL
    );
    CREATE TABLE IF NOT EXISTS users(
      id INTEGER PRIMARY KEY AUTOINCREMENT, hospital_id TEXT NOT NULL, username TEXT UNIQUE NOT NULL,
      password_hash TEXT NOT NULL, role TEXT DEFAULT 'admin', created_at TEXT NOT NULL
    );
    CREATE TABLE IF NOT EXISTS patients(
      id INTEGER PRIMARY KEY AUTOINCREMENT, patient_id TEXT UNIQUE NOT NULL, hospital_id TEXT NOT NULL,
      name TEXT NOT NULL, age INTEGER, gender TEXT, phone TEXT, blood_group TEXT, address TEXT,
      emergency_contact TEXT, created_at TEXT NOT NULL
    );
    CREATE TABLE IF NOT EXISTS allergies(
      id INTEGER PRIMARY KEY AUTOINCREMENT, patient_id TEXT NOT NULL, allergen TEXT NOT NULL,
      reaction TEXT, severity TEXT DEFAULT 'Moderate', created_at TEXT NOT NULL
    );
    CREATE TABLE IF NOT EXISTS medicines(
      id INTEGER PRIMARY KEY AUTOINCREMENT, medicine_id TEXT UNIQUE NOT NULL, name TEXT NOT NULL,
      generic_name TEXT, strength TEXT, dosage_form TEXT, route TEXT, category TEXT, stock INTEGER DEFAULT 0,
      allergy_class TEXT, created_at TEXT NOT NULL
    );
    CREATE TABLE IF NOT EXISTS drug_interactions(
      id INTEGER PRIMARY KEY AUTOINCREMENT, medicine_a TEXT NOT NULL, medicine_b TEXT NOT NULL,
      severity TEXT NOT NULL, description TEXT, recommendation TEXT
    );
    CREATE TABLE IF NOT EXISTS prescriptions(
      id INTEGER PRIMARY KEY AUTOINCREMENT, prescription_id TEXT UNIQUE NOT NULL, patient_id TEXT NOT NULL,
      doctor_name TEXT, diagnosis TEXT, notes TEXT, status TEXT DEFAULT 'Issued', created_at TEXT NOT NULL
    );
    CREATE TABLE IF NOT EXISTS prescription_medicines(
      id INTEGER PRIMARY KEY AUTOINCREMENT, prescription_id TEXT NOT NULL, medicine_id TEXT,
      medicine_name TEXT NOT NULL, dose TEXT, frequency TEXT, duration TEXT, route TEXT, instructions TEXT
    );
    CREATE TABLE IF NOT EXISTS safety_results(
      id INTEGER PRIMARY KEY AUTOINCREMENT, prescription_id TEXT NOT NULL, patient_id TEXT NOT NULL,
      overall_status TEXT NOT NULL, findings_json TEXT NOT NULL, created_at TEXT NOT NULL
    );
    """)
    c.commit(); c.close()

def hpw(p): return hashlib.sha256(p.encode()).hexdigest()

def seed_admin():
    """Create the built-in platform administrator once.
    Hospitals still create their own separate accounts via registration.
    """
    c = db()
    now = datetime.now().isoformat()
    c.execute("""INSERT OR IGNORE INTO hospitals(hospital_id,hospital_name,email,phone,address,created_at)
                 VALUES(?,?,?,?,?,?)""",
              ("MEDSAFE-ADMIN", "MedSafe AI Administration", "admin@medsafe.local", "", "Platform administration", now))
    c.execute("""INSERT OR IGNORE INTO users(hospital_id,username,password_hash,role,created_at)
                 VALUES(?,?,?,?,?)""",
              ("MEDSAFE-ADMIN", "admin", hpw("admin123"), "super_admin", now))
    c.commit(); c.close()
def next_id(prefix, table, column, width=6):
    c=db(); n=c.execute(f"SELECT COUNT(*) n FROM {table}").fetchone()["n"]+1; c.close()
    return f"{prefix}-{n:0{width}d}"

def logged():
    return bool(session.get("logged_in"))

# ---------- XLSX IMPORT ----------
def clean_key(v):
    return re.sub(r"[^a-z0-9]+","_",str(v or "").strip().lower()).strip("_")

def val(row, *names):
    normalized={clean_key(k): row[k] for k in row.keys()}
    for name in names:
        if clean_key(name) in normalized:
            x=normalized[clean_key(name)]
            if x is not None and str(x).strip()!="": return str(x).strip()
    return ""

def import_xlsx(path):
    """Import the MedSafe workbook using its sheet relationships.
    Supports the supplied workbook structure: Drug Master, Indications,
    Product Barcode, Dosing, Renal Adjustment, Interactions,
    Contraindication, Allergies, Adverse Effect and Alternatives.
    """
    try:
        import openpyxl
    except ImportError:
        return {"success":False,"message":"openpyxl is missing. Run: pip install -r requirements.txt"}

    wb=openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheets={ws.title.strip().lower(): ws for ws in wb.worksheets}

    def read_sheet(ws):
        if not ws: return []
        rows=ws.iter_rows(values_only=True)
        try: headers=next(rows)
        except StopIteration: return []
        headers=[str(h or '').strip() or f'column_{i}' for i,h in enumerate(headers)]
        return [dict(zip(headers, values)) for values in rows]

    c=db(); count=0
    drug_by_id={}; name_by_id={}

    # 1) Drug Master is the authoritative medicine list.
    drug_rows=read_sheet(sheets.get('drug master'))
    for row in drug_rows:
        did=val(row,'drug_id','medicine_id','id')
        generic=val(row,'generic_name','generic name')
        brand=val(row,'brand_name','brand name')
        name=generic or brand
        if not did or not name: continue
        drug_by_id[did]=row
        name_by_id[did]=name

    # Allergy class comes from the workbook's Allergies sheet.
    allergy_rows=read_sheet(sheets.get('allergies'))
    allergy_class_by_id={}
    for row in allergy_rows:
        did=val(row,'drug_id')
        cls=val(row,'drug_class')
        if did and cls: allergy_class_by_id[did]=cls

    for did,row in drug_by_id.items():
        generic=val(row,'generic_name','generic name')
        brand=val(row,'brand_name','brand name')
        name=generic or brand
        category=val(row,'therapeutic_class','therapeutic class') or val(row,'drug_class','drug class')
        c.execute("""INSERT INTO medicines(medicine_id,name,generic_name,strength,dosage_form,route,category,stock,allergy_class,created_at)
                     VALUES(?,?,?,?,?,?,?,?,?,?)
                     ON CONFLICT(medicine_id) DO UPDATE SET
                     name=excluded.name,generic_name=excluded.generic_name,strength=excluded.strength,
                     dosage_form=excluded.dosage_form,route=excluded.route,category=excluded.category,
                     allergy_class=excluded.allergy_class""",
                  (did,name,generic,'',val(row,'dosage_form','dosage form'),val(row,'route'),category,0,
                   allergy_class_by_id.get(did,''),datetime.now().isoformat()))
        count+=1

    # 2) Interactions are a separate relational sheet keyed by Drug ID.
    interactions=0
    for row in read_sheet(sheets.get('interactions')):
        a_id=val(row,'drug_a_id'); b_id=val(row,'drug_b_id')
        a=name_by_id.get(a_id,a_id); b=name_by_id.get(b_id,b_id)
        sev=val(row,'severity')
        if not a or not b or not sev: continue
        desc=val(row,'clinical_effect','description','interaction')
        mech=val(row,'mechanism')
        mgmt=val(row,'management','recommendation','action')
        description=(f"{mech}. {desc}" if mech and desc else mech or desc)
        c.execute("""INSERT INTO drug_interactions(medicine_a,medicine_b,severity,description,recommendation)
                     SELECT ?,?,?,?,? WHERE NOT EXISTS(
                     SELECT 1 FROM drug_interactions WHERE
                     (lower(medicine_a)=lower(?) AND lower(medicine_b)=lower(?)) OR
                     (lower(medicine_a)=lower(?) AND lower(medicine_b)=lower(?)))""",
                  (a,b,sev,description,mgmt,a,b,b,a))
        interactions+=1

    # 3) Keep workbook-backed safety metadata in a JSON sidecar for future checks/UI.
    metadata={
      'indications': read_sheet(sheets.get('indications')),
      'dosing': read_sheet(sheets.get('dosing')),
      'renal_adjustment': read_sheet(sheets.get('renal adjustment')),
      'contraindications': read_sheet(sheets.get('contraindication')) or read_sheet(sheets.get('contraindication ')),
      'allergies': allergy_rows,
      'adverse_effects': read_sheet(sheets.get('adverse effect')),
      'alternatives': read_sheet(sheets.get('alternatives')),
      'product_barcode': read_sheet(sheets.get('product barcode')),
    }
    meta_path=BASE_DIR/'database'/'medsafe_master_metadata.json'
    with open(meta_path,'w',encoding='utf-8') as f: json.dump(metadata,f,ensure_ascii=False,default=str)

    c.commit(); c.close(); wb.close()
    return {"success":True,"medicines":count,"interactions":interactions,"sheets":len(wb.sheetnames) if hasattr(wb,'sheetnames') else len(sheets),"file":str(path)}

def auto_import():
    for p in XLSX_DEFAULTS:
        if p.exists():
            return import_xlsx(p)
    return None

# ---------- PAGES ----------
@app.route("/")
def login_page():
    return redirect("/dashboard") if logged() else render_template("login.html")

@app.route("/register")
def register_page(): return render_template("register.html")

@app.route("/dashboard")
def dashboard():
    return render_template("dashboard.html") if logged() else redirect("/")

@app.route("/patients")
def patients_page():
    return render_template("patients.html") if logged() else redirect("/")

@app.route("/prescription")
def prescription_page():
    return render_template("prescription.html") if logged() else redirect("/")

@app.route("/safety-analysis")
def safety_page():
    return render_template("safety-analysis.html") if logged() else redirect("/")

@app.route("/medicine-verification")
def verify_page():
    return render_template("medicine-verification.html") if logged() else redirect("/")

@app.route("/prescription-verification")
def prescription_verify_page():
    return render_template("prescription-verification.html") if logged() else redirect("/")

# ---------- AUTH ----------
@app.post("/api/register-hospital")
def register_hospital():
    d=request.get_json() or {}
    required=("hospital_name","email","username","password")
    if any(not str(d.get(k,"")).strip() for k in required):
        return jsonify(success=False,message="Please complete all required fields."),400
    c=db()
    try:
        hid=next_id("HSP","hospitals","hospital_id",4)
        c.execute("INSERT INTO hospitals(hospital_id,hospital_name,email,phone,address,created_at) VALUES(?,?,?,?,?,?)",
                  (hid,d["hospital_name"].strip(),d["email"].strip().lower(),d.get("phone",""),d.get("address",""),datetime.now().isoformat()))
        c.execute("INSERT INTO users(hospital_id,username,password_hash,role,created_at) VALUES(?,?,?,?,?)",
                  (hid,d["username"].strip(),hpw(d["password"]),"admin",datetime.now().isoformat()))
        c.commit()
        return jsonify(success=True,message="Hospital registered successfully.",hospital_id=hid)
    except sqlite3.IntegrityError:
        c.rollback(); return jsonify(success=False,message="Hospital email or username already exists."),409
    finally: c.close()

@app.post("/api/login")
def login():
    d=request.get_json() or {}; u=str(d.get("username","")).strip(); p=str(d.get("password",""))
    c=db(); row=c.execute("""SELECT u.*,h.hospital_name FROM users u JOIN hospitals h ON h.hospital_id=u.hospital_id
                             WHERE u.username=? AND u.password_hash=?""",(u,hpw(p))).fetchone(); c.close()
    if not row: return jsonify(success=False,message="Invalid username or password."),401
    session.update(logged_in=True,hospital_id=row["hospital_id"],hospital_name=row["hospital_name"],username=row["username"],role=row["role"])
    return jsonify(success=True)

@app.post("/api/logout")
def logout(): session.clear(); return jsonify(success=True)

# ---------- DASHBOARD ----------
@app.get("/api/dashboard")
def dashboard_api():
    if not logged(): return jsonify(success=False),401
    c=db()
    hid=session.get("hospital_id")
    is_admin=session.get("role") == "super_admin"
    if is_admin:
        stats={
          "patients":c.execute("SELECT COUNT(*) n FROM patients").fetchone()["n"],
          "prescriptions":c.execute("SELECT COUNT(*) n FROM prescriptions").fetchone()["n"],
          "medicines":c.execute("SELECT COUNT(*) n FROM medicines").fetchone()["n"],
          "safety_checks":c.execute("SELECT COUNT(*) n FROM safety_results").fetchone()["n"]
        }
        recent=c.execute("""SELECT p.prescription_id,p.patient_id,p.created_at,p.status,pt.name patient_name
                            FROM prescriptions p LEFT JOIN patients pt ON pt.patient_id=p.patient_id
                            ORDER BY p.id DESC LIMIT 8""").fetchall()
    else:
        stats={
          "patients":c.execute("SELECT COUNT(*) n FROM patients WHERE hospital_id=?",(hid,)).fetchone()["n"],
          "prescriptions":c.execute("""SELECT COUNT(*) n FROM prescriptions p JOIN patients pt ON pt.patient_id=p.patient_id
                                      WHERE pt.hospital_id=?""",(hid,)).fetchone()["n"],
          "medicines":c.execute("SELECT COUNT(*) n FROM medicines").fetchone()["n"],
          "safety_checks":c.execute("""SELECT COUNT(*) n FROM safety_results s JOIN patients pt ON pt.patient_id=s.patient_id
                                       WHERE pt.hospital_id=?""",(hid,)).fetchone()["n"]
        }
        recent=c.execute("""SELECT p.prescription_id,p.patient_id,p.created_at,p.status,pt.name patient_name
                            FROM prescriptions p JOIN patients pt ON pt.patient_id=p.patient_id
                            WHERE pt.hospital_id=? ORDER BY p.id DESC LIMIT 8""",(hid,)).fetchall()
    c.close()
    return jsonify(success=True,stats=stats,recent=[dict(x) for x in recent])

# ---------- PATIENTS ----------
@app.post("/api/register-patient")
def register_patient():
    if not logged(): return jsonify(success=False),401
    d=request.get_json() or {}
    if not str(d.get("name","")).strip(): return jsonify(success=False,message="Patient name is required."),400
    c=db(); pid=next_id("PT","patients","patient_id")
    try:
        c.execute("""INSERT INTO patients(patient_id,hospital_id,name,age,gender,phone,blood_group,address,emergency_contact,created_at)
                     VALUES(?,?,?,?,?,?,?,?,?,?)""",
                  (pid,session.get("hospital_id"),d["name"].strip(),int(d["age"]) if str(d.get("age","")).isdigit() else None,
                   d.get("gender",""),d.get("phone",""),d.get("blood_group",""),d.get("address",""),d.get("emergency_contact",""),datetime.now().isoformat()))
        for a in d.get("allergies",[]):
            if str(a.get("allergen","")).strip():
                c.execute("INSERT INTO allergies(patient_id,allergen,reaction,severity,created_at) VALUES(?,?,?,?,?)",
                          (pid,a["allergen"].strip(),a.get("reaction",""),a.get("severity","Moderate"),datetime.now().isoformat()))
        c.commit(); return jsonify(success=True,patient_id=pid)
    finally: c.close()

@app.get("/api/patients")
def patients_api():
    if not logged(): return jsonify(success=False),401
    c=db()
    rows=c.execute("""SELECT pt.*,
                      (SELECT COUNT(*) FROM prescriptions p WHERE p.patient_id=pt.patient_id) prescription_count,
                      (SELECT COUNT(*) FROM allergies a WHERE a.patient_id=pt.patient_id) allergy_count
                      FROM patients pt WHERE pt.hospital_id=? ORDER BY pt.id DESC""",(session.get("hospital_id"),)).fetchall()
    c.close(); return jsonify(success=True,patients=[dict(x) for x in rows])

@app.get("/api/patient/<pid>")
def patient_api(pid):
    if not logged(): return jsonify(success=False),401
    c=db(); p=c.execute("SELECT * FROM patients WHERE patient_id=? AND hospital_id=?",(pid,session.get("hospital_id"))).fetchone()
    if not p: c.close(); return jsonify(success=False,message="Patient not found."),404
    a=c.execute("SELECT * FROM allergies WHERE patient_id=? ORDER BY id DESC",(pid,)).fetchall()
    r=c.execute("SELECT * FROM prescriptions WHERE patient_id=? ORDER BY id DESC",(pid,)).fetchall()
    c.close(); return jsonify(success=True,patient=dict(p),allergies=[dict(x) for x in a],prescriptions=[dict(x) for x in r])

# ---------- MEDICINES ----------
@app.get("/api/medicines/search")
def medicine_search():
    if not logged(): return jsonify(success=False),401
    q=request.args.get("q","").strip()
    c=db()
    if q:
        rows=c.execute("""SELECT * FROM medicines WHERE name LIKE ? OR generic_name LIKE ? OR medicine_id LIKE ?
                          ORDER BY name LIMIT 30""",(f"%{q}%",f"%{q}%",f"%{q}%")).fetchall()
    else: rows=c.execute("SELECT * FROM medicines ORDER BY name LIMIT 30").fetchall()
    c.close(); return jsonify(success=True,medicines=[dict(x) for x in rows])

@app.post("/api/verify-medicine")
def verify_medicine():
    if not logged(): return jsonify(success=False),401
    term=str((request.get_json() or {}).get("medicine","")).strip()
    c=db(); row=c.execute("""SELECT * FROM medicines WHERE medicine_id=? OR lower(name)=lower(?) OR lower(generic_name)=lower(?) LIMIT 1""",
                          (term,term,term)).fetchone(); c.close()
    if not row: return jsonify(success=True,verified=False,message="Medicine not found in the built-in medicine master.")
    return jsonify(success=True,verified=True,medicine=dict(row))

@app.get("/api/prescriptions/search")
def prescription_search():
    if not logged(): return jsonify(success=False),401
    q=request.args.get("q","").strip().lower()
    c=db(); hid=session.get("hospital_id")
    rows=c.execute("""SELECT p.prescription_id,p.patient_id,p.doctor_name,p.diagnosis,p.status,p.created_at,pt.name patient_name
                     FROM prescriptions p JOIN patients pt ON pt.patient_id=p.patient_id
                     WHERE pt.hospital_id=? AND (lower(p.prescription_id) LIKE ? OR lower(p.patient_id) LIKE ? OR lower(pt.name) LIKE ? OR lower(coalesce(p.doctor_name,'')) LIKE ?)
                     ORDER BY p.id DESC LIMIT 12""",(hid,f"%{q}%",f"%{q}%",f"%{q}%",f"%{q}%")).fetchall() if q else c.execute("""SELECT p.prescription_id,p.patient_id,p.doctor_name,p.diagnosis,p.status,p.created_at,pt.name patient_name FROM prescriptions p JOIN patients pt ON pt.patient_id=p.patient_id WHERE pt.hospital_id=? ORDER BY p.id DESC LIMIT 12""",(hid,)).fetchall()
    c.close(); return jsonify(success=True,prescriptions=[dict(x) for x in rows])

@app.get("/api/prescription/<rx_id>")
def prescription_detail(rx_id):
    if not logged(): return jsonify(success=False),401
    c=db(); hid=session.get("hospital_id")
    row=c.execute("""SELECT p.*,pt.name patient_name,pt.hospital_id FROM prescriptions p JOIN patients pt ON pt.patient_id=p.patient_id
                     WHERE p.prescription_id=? AND pt.hospital_id=?""",(rx_id,hid)).fetchone()
    if not row: c.close(); return jsonify(success=False,message="Prescription not found in this hospital workspace."),404
    meds=c.execute("SELECT * FROM prescription_medicines WHERE prescription_id=? ORDER BY id",(rx_id,)).fetchall()
    c.close(); return jsonify(success=True,prescription=dict(row),medicines=[dict(x) for x in meds])

# ---------- PRESCRIPTIONS ----------
@app.post("/api/create-prescription")
def create_prescription():
    if not logged(): return jsonify(success=False),401
    d=request.get_json() or {}; pid=str(d.get("patient_id","")).strip(); meds=d.get("medicines",[])
    if not pid or not meds: return jsonify(success=False,message="Patient and at least one medicine are required."),400
    c=db(); patient=c.execute("SELECT * FROM patients WHERE patient_id=? AND hospital_id=?",(pid,session.get("hospital_id"))).fetchone()
    if not patient: c.close(); return jsonify(success=False,message="Patient not found."),404
    rx=next_id("RX","prescriptions","prescription_id",6)
    c.execute("""INSERT INTO prescriptions(prescription_id,patient_id,doctor_name,diagnosis,notes,status,created_at)
                 VALUES(?,?,?,?,?,?,?)""",(rx,pid,d.get("doctor_name",""),d.get("diagnosis",""),d.get("notes",""),"Issued",datetime.now().isoformat()))
    for m in meds:
        c.execute("""INSERT INTO prescription_medicines(prescription_id,medicine_id,medicine_name,dose,frequency,duration,route,instructions)
                     VALUES(?,?,?,?,?,?,?,?)""",(rx,m.get("medicine_id",""),m.get("medicine_name",""),m.get("dose",""),m.get("frequency",""),
                                                m.get("duration",""),m.get("route",""),m.get("instructions","")))
    c.commit(); c.close(); return jsonify(success=True,prescription_id=rx)

# ---------- SAFETY ----------
def allergy_match(allergen, med):
    a=allergen.lower().strip(); m=med.lower().strip()
    if not a or not m: return False
    groups={
      "penicillin":["amoxicillin","ampicillin","penicillin","amoxicillin-clavulanate"],
      "beta lactam":["amoxicillin","ampicillin","penicillin"],
      "nsaid":["ibuprofen","aspirin","diclofenac","naproxen"],
      "sulfa":["sulfamethoxazole","sulfonamide"],
    }
    for group,items in groups.items():
        if group in a and any(x in m for x in items): return True
    return a in m or m in a

@app.post("/api/safety-analysis")
def safety():
    if not logged(): return jsonify(success=False),401
    d=request.get_json() or {}; pid=str(d.get("patient_id","")).strip(); rx=str(d.get("prescription_id","")).strip()
    if not pid or not rx: return jsonify(success=False,message="Patient ID and Prescription ID are required."),400
    c=db()
    p=c.execute("SELECT * FROM patients WHERE patient_id=? AND hospital_id=?",(pid,session.get("hospital_id"))).fetchone()
    pr=c.execute("SELECT * FROM prescriptions WHERE prescription_id=? AND patient_id=?",(rx,pid)).fetchone()
    if not p or not pr: c.close(); return jsonify(success=False,message="Patient or prescription not found."),404
    allergies=c.execute("SELECT * FROM allergies WHERE patient_id=?",(pid,)).fetchall()
    meds=c.execute("SELECT * FROM prescription_medicines WHERE prescription_id=?",(rx,)).fetchall()
    findings=[]
    for m in meds:
        for a in allergies:
            if allergy_match(a["allergen"],m["medicine_name"]):
                findings.append({"type":"ALLERGY","severity":"RED","medicine":m["medicine_name"],
                                 "message":f"Possible allergy conflict: recorded allergy to {a['allergen']}.",
                                 "recommendation":"Hold and clinically review before administration."})
    names=[m["medicine_name"] for m in meds]
    for i in range(len(names)):
        for j in range(i+1,len(names)):
            x=c.execute("""SELECT * FROM drug_interactions WHERE
                (lower(medicine_a)=lower(?) AND lower(medicine_b)=lower(?)) OR
                (lower(medicine_a)=lower(?) AND lower(medicine_b)=lower(?)) LIMIT 1""",(names[i],names[j],names[j],names[i])).fetchone()
            if x:
                s=x["severity"].upper()
                findings.append({"type":"DRUG INTERACTION","severity":"RED" if s=="HIGH" else "YELLOW",
                                 "medicine":f"{names[i]} + {names[j]}","message":x["description"] or "Interaction found in medicine master.",
                                 "recommendation":x["recommendation"] or "Review combination before dispensing."})
    overall="RED" if any(x["severity"]=="RED" for x in findings) else "YELLOW" if findings else "GREEN"
    c.execute("INSERT INTO safety_results(prescription_id,patient_id,overall_status,findings_json,created_at) VALUES(?,?,?,?,?)",
              (rx,pid,overall,json.dumps(findings),datetime.now().isoformat()))
    c.commit(); c.close()
    return jsonify(success=True,overall_status=overall,findings=findings,
                    message={"RED":"Critical safety issue detected.","YELLOW":"Potential safety issue requires review.","GREEN":"No rule-based conflicts detected."}[overall])

@app.get("/api/health")
def health():
    return jsonify(status="ok")

if __name__=="__main__":
    init_db()
    seed_admin()
    app.run(host="0.0.0.0",port=int(os.environ.get("PORT",5000)),debug=False)
